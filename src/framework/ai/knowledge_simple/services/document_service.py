# src/framework/ai/knowledge_simple/services/document_service.py
import os
import uuid
import re
from typing import List, Tuple, Optional
from sqlalchemy.orm import Session
from sqlalchemy import or_, func, text, bindparam
from ..database import KnowledgeDocument, KnowledgeChunk
from .embedding_service import EmbeddingService
from .qdrant_service import QdrantService
# 导入日志
import logging
logger = logging.getLogger(__name__)

class DocumentService:
    """文档处理服务"""

    def __init__(self):
        self.embedding_service = EmbeddingService()
        self.qdrant = QdrantService()
        self.chunk_size = 500  # 每块字符数
        self.chunk_overlap = 50  # 重叠字符数
        self.use_embedding = self.embedding_service.enabled


    def extract_text_from_file(self,file_path: str, file_type: str) -> str:
        """从文件中提取文本"""
        try:
            if file_type == 'txt':
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read() 
            elif file_type == 'pdf':
                try:
                    import pdfplumber 
                    with pdfplumber.open(file_path) as pdf:
                        text = ''
                        for page in pdf.pages:
                            page_text = page.extract_text() or ''
                            text += page_text + "\n"
                    return text
                except ImportError:
                    return "⚠️ 请安装 pdfplumber: pip install pdfplumber"
            elif file_type == 'docx':
                try:
                    from docx import Document
                    doc = Document(file_path)
                    text = '\n'.join([para.text for para in doc.paragraphs])
                    return text
                except ImportError:
                    return "⚠️ 请安装 python-docx: pip install python-docx"
                except Exception as e:
                    return f"DOCX解析失败: {e}"
            else:
                return f"不支持的文件类型: {file_type}"
        except Exception as e:
            return f"提取文本失败: {e}"


    def chunk_text(self, text: str) -> List[str]:
        """
            将文本分块,滑动窗口算法。
            有点难啊。
        """
        if not text or not text.strip():
            return []
        # 清理文本
        text = text.strip()
        chunks = []
        start = 0
        text_length = len(text)

        while start < text_length:
            end = start + self.chunk_size
            if end < text_length:
                # 尝试在句号、问号等处分段
                for sep in ['。', '！', '？', '\n\n', '\n', '.', '!', '?']:
                    pos = text.rfind(sep, start, end)
                    if pos > start:
                        end = pos + len(sep)
                        break
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)

            start = end - self.chunk_overlap if end < text_length else end
        return chunks

    # ========== 3. 保存文档：PG 文本 + Qdrant 向量 ==========
    def save_document(
            self, 
            db: Session, 
            title: str, 
            filename: str, 
            file_type: str, 
            content: str
    ) -> KnowledgeDocument:
        """保存文档并生成向量"""
        # 1. 创建文档记录
        doc = KnowledgeDocument(
            title=title or filename,
            filename=filename,
            file_type=file_type,
            content=content[:10000] # 只保存前10000字符，避免过大
        )
        db.add(doc)
        db.flush()  # 拿到 doc.id
        # 2. 分块 + 生成向量
        chunks = self.chunk_text(content)
        pg_chunks: List[KnowledgeChunk] = []
        qdrant_points: List[dict] = []
        # 3. 为每块生成向量并保存
        for idx, chunk_text in enumerate(chunks):
            chunk_id = uuid.uuid4() # ★ 提前生成 UUID，两边共用
            embedding = None
            # if self.embedding_service:
            if self.use_embedding:
                embedding = self.embedding_service.get_embedding(chunk_text)
                print(f"📊 分块 {idx} 向量维度: {len(embedding) if embedding else 'None'}")   # ← 加这行
          
            pg_chunks.append(KnowledgeChunk(
                id = chunk_id,
                chunk_index=idx,
                document_id=doc.id,
                content=chunk_text,
            ))
            if embedding:
                qdrant_points.append({
                    "chunk_id":str(chunk_id),
                    "vector" : embedding,
                    "payload": {
                        "document_id": str(doc.id),
                        "document_title": doc.title,
                        "chunk_index" :idx,
                        "content" : chunk_text,
                    },
                })
        db.add_all(pg_chunks)
        doc.chunk_count = len(pg_chunks)
        db.commit()
        db.refresh(doc)
        # 4) 批量写 Qdrant
        if qdrant_points:
            try:
                n = self.qdrant.upsert_chunks(qdrant_points)
                logger.info(f"✅ '{doc.title}':PG{len(pg_chunks)}块 / Qdrant {n} 向量")
            except Exception as e:
                logger.info(f" Qdrant 写入失败: {e}")
        return doc

    # ========== 4. 搜索：改走 Qdrant ==========
    def search_similar(
            self,
            db: Session, 
            query: str,
            top_k: int = 5,
            threshold: float = 0.5
    ) -> List[Tuple[KnowledgeChunk, float]]:
        """搜索相似文档块"""
        # 无 embedding 服务 → 降级到文本搜索
        if not self.use_embedding:
            return self._search_by_text(db, query, top_k)
        query_embedding = None
        query_vector = self.embedding_service.get_embedding(query)
        if not query_vector:
            # 降级到文本搜索
            return self._search_by_text(db, query, top_k)
        
        # Qdrant 检索
        try:
           hits = self.qdrant.search(query_vector, top_k=top_k, threshold=threshold)
        except Exception as e:
            print(f"向量搜索失败，降级到文本搜索: {e}") 
            return self._search_by_text(db, query, top_k)
        if not hits:
            return []
        
        # 回 PG 拿完整 chunk（补 created_at 等）
        ids = [h["chunk_id"] for h in hits]
        chunks = db.query(KnowledgeChunk).filter(KnowledgeChunk.id.in_(ids)).all()
        chunk_map = {str(c.id): c for c in chunks}
        return [
            (chunk_map[h["chunk_id"]], h["score"] )for h in hits if h["chunk_id"] in chunk_map
        ]

    # ========== 5. 文本降级搜索（保留原样） ==========
    def _search_by_text(
            self,
            db:Session,
            query: str, 
            top_k: int =5
    ) -> List[Tuple[KnowledgeChunk, float]]:
        """基于文本的搜索（降级方案）"""
        try:
            # 使用 PostgreSQL 全文搜索或 LIKE
            search_pattern = f"%{query}%"
            results = db.query(KnowledgeChunk).filter(KnowledgeChunk.content.ilike(search_pattern)).limit(top_k).all() #数据表查到的。
            # 计算简单的相关性分数（关键词匹配度）
            query_words = set(query.lower().split()) #请求过来的。
            scored_results = []

            # 数据库查询到相似的内容拿出来
            for chunk in results:
                chunk_words = set(chunk.content.lower().split())
                if chunk_words:
                    overlap = len(chunk_words & query_words) / len(query_words) # 重叠词数量  ÷  查询词总数 = 相似度分数
                    scored_results.append((chunk, overlap)) # 存储 (分块, 相似度)

            scored_results.sort(key=lambda x:x[1], reverse=True)
            return scored_results
        except Exception as e:
            print(f"{e}")


    def get_document_by_id(self,db:Session,doc_id: str) -> Optional[KnowledgeDocument]:
        """根据ID获取文档"""
        try:
            return db.query(KnowledgeDocument).filter(KnowledgeDocument.id == uuid.UUID(doc_id)).first()
        except ValueError:
            return None

    # ========== 6. 删除文档：PG + Qdrant 一起删 ==========
    def delete_document(self, db: Session, doc_id: str) -> bool:
        try:
            doc_uuid = uuid.UUID(doc_id)
        except ValueError:
            return False
        doc = db.query(KnowledgeDocument).filter(
            KnowledgeDocument.id == doc_uuid
        ).first()
        if not doc:
            return False

        try:
            self.qdrant.delete_by_document(doc_id)
        except Exception as e:
            logger.exception(f"❌ Qdrant 删除失败: {e}")

        db.delete(doc) # cascade 会删 chunks
        db.commit()
        return True

